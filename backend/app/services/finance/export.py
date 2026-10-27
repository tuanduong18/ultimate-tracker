"""The Transactions tab as a spreadsheet: the rows on screen, as .xlsx.

The tab searches and filters in the browser, so it already knows which rows
someone is looking at; it sends their ids rather than its filters. That way
the filters live in one place, and the file cannot disagree with the screen
because a filter was written twice and one copy drifted.

Built here rather than in the browser so amounts go into Excel as the exact
decimals stored, not as floats that passed through JavaScript, and so the
converted column can use the server's exchange rates.
"""

import uuid
from collections.abc import Sequence
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.currencies import minor_units
from app.models.finance import Category, Expense
from app.services import exchange_rates
from app.services.finance.categories import list_categories
from app.services.finance.expenses import list_expenses_by_id
from app.services.finance.summaries import UNCATEGORIZED_NAME

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The columns an expense is written as, in order. The export leads with them
# and the import template is exactly them, so an exported file imports as it is.
COLUMNS = ("Date", "Description", "Category", "Amount", "Currency")

_BOLD = Font(bold=True)
_DATE_FORMAT = "yyyy-mm-dd"


def _money_format(currency: str) -> str:
    """Thousands separators and exactly the decimals the currency has: VND 0, USD 2."""
    places = minor_units(currency)
    return "#,##0" + ("." + "0" * places if places else "")


def _text(cell: Cell, value: str) -> None:
    """Store a string as a string, whatever it starts with.

    openpyxl turns any string beginning with "=" into a formula, so a
    description someone typed as "=HYPERLINK(...)" would become a live link in
    the exported file. A typed string cell is never evaluated by Excel.
    """
    cell.value = value
    cell.data_type = "s"


def build_workbook(
    expenses: Sequence[Expense],
    categories: Sequence[Category],
    *,
    filters: Sequence[str],
    display_currency: str,
    rates: dict[str, Decimal],
    exported_at: datetime,
) -> bytes:
    """The workbook as bytes: a Transactions sheet and a sheet saying what it holds.

    The converted column is what makes the sheet summable across currencies;
    a row whose rate is missing leaves that cell blank rather than writing 0,
    so a SUM over the column comes up visibly short instead of silently wrong.
    """
    names = {c.id: c.name for c in categories}
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Transactions"

    converted_header = f"Amount ({display_currency})"
    headers = [*COLUMNS, converted_header]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _BOLD

    for row, expense in enumerate(expenses, start=2):
        sheet.cell(row=row, column=1, value=expense.spent_on).number_format = _DATE_FORMAT
        _text(sheet.cell(row=row, column=2), expense.description or "")
        category = names.get(expense.category_id) if expense.category_id else None
        _text(sheet.cell(row=row, column=3), category or UNCATEGORIZED_NAME)
        amount = sheet.cell(row=row, column=4, value=expense.amount)
        amount.number_format = _money_format(expense.currency)
        _text(sheet.cell(row=row, column=5), expense.currency)
        try:
            value = exchange_rates.convert(
                expense.amount, expense.currency, display_currency, rates
            )
        except ValueError:
            continue
        converted = sheet.cell(row=row, column=6, value=value)
        converted.number_format = _money_format(display_currency)

    # Headers stay in view while scrolling, and each gets a sort/filter arrow.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    widths = [12, 40, 18, 14, 10, max(16, len(converted_header) + 2)]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    about = book.create_sheet("About this export")
    about.append(["Exported at", exported_at.strftime("%Y-%m-%d %H:%M %Z").strip()])
    about.append(["Transactions", len(expenses)])
    about.append([converted_header, f"Converted at the exchange rates of {exported_at:%Y-%m-%d}"])
    about.append([])
    about.append(["Filters"])
    for line in filters or ["None — every transaction"]:
        _text(about.cell(row=about.max_row + 1, column=1), line)
    for row in (1, 2, 3, 5):
        about.cell(row=row, column=1).font = _BOLD
    about.column_dimensions["A"].width = 24
    about.column_dimensions["B"].width = 48

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def export_expenses(
    db: AsyncSession,
    user_id: uuid.UUID,
    expense_ids: Sequence[uuid.UUID],
    *,
    filters: Sequence[str],
    display_currency: str,
    exported_at: datetime,
) -> bytes:
    """The user's expenses among ``expense_ids``, newest first, as an .xlsx file.

    ``filters`` is the client's own description of what produced the list
    ("Search: taxi"), written verbatim onto the About sheet as text.
    """
    expenses = await list_expenses_by_id(db, user_id, expense_ids)
    categories = await list_categories(db, user_id)
    # get_rates never raises; a currency it does not cover costs that row's
    # converted cell (see build_workbook), not the whole export.
    rates = await exchange_rates.get_rates(db)
    return build_workbook(
        expenses,
        categories,
        filters=filters,
        display_currency=display_currency,
        rates=rates,
        exported_at=exported_at,
    )
