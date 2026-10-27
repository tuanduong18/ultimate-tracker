"""Expenses from a spreadsheet: the template, reading an upload, and saving it.

The Transactions tab drives it in three steps:

1. ``import_template`` — an .xlsx whose Expenses sheet has exactly the columns
   the importer reads (the five the export leads with, in the same order, so an
   exported file imports as it is), drop-downs of the user's own categories and
   the supported currencies, and a sheet that explains each column.
2. ``preview_import`` — read an uploaded file and say what importing it would
   do: the rows ready to go, the rows that can't go and why, which look like
   expenses already recorded, which categories would be created. Saves nothing.
3. ``import_expenses`` — the same reading, then one transaction: the ready rows
   (less the likely duplicates, unless asked to keep them) and any categories
   they name that don't exist yet.

The browser sends the file twice rather than the server keeping it between the
two calls, so nothing is stored, and what is saved is decided by the code that
drew the preview, against the user's data as it is when they confirm.

Fixed columns rather than a column-mapping step: the template is where a bank
statement gets pasted, and a file in the template's shape needs no questions.
Read on the server, like the export is written there, so amounts arrive as the
decimals in the file rather than as JavaScript floats.
"""

import csv
import math
import re
import uuid
import zipfile
from collections import Counter
from collections.abc import Iterable, Sequence
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.currencies import minor_units
from app.models.finance import Category, Expense
from app.schemas.finance import check_amount_scale
from app.services import exchange_rates
from app.services.finance.categories import list_categories
from app.services.finance.errors import CategoryNameTakenError, ImportFileError
from app.services.finance.export import _BOLD, _DATE_FORMAT, COLUMNS, _text

# Bigger than any honest file of MAX_IMPORT_ROWS rows, which is a few hundred KB.
MAX_IMPORT_BYTES = 5 * 1024 * 1024
# Years of daily spending, and a bound on what one request writes.
MAX_IMPORT_ROWS = 5_000
# What an .xlsx may unpack to. A real one of MAX_IMPORT_ROWS rows is a few MB;
# far beyond that is a zip bomb, refused before any of it is parsed.
_MAX_UNPACKED_BYTES = 100 * 1024 * 1024

# Rows of the template that carry the date format and the drop-downs.
_TEMPLATE_ROWS = 1_000

_DATE, _DESCRIPTION, _CATEGORY, _AMOUNT, _CURRENCY = COLUMNS
_REQUIRED = (_DATE, _AMOUNT)
# Column names are matched however they are capitalised or padded.
_BY_KEY = {column.casefold(): column for column in COLUMNS}

# Read as "no category", the way the app labels one.
_UNCATEGORISED = frozenset({"uncategorised", "uncategorized"})
_DESCRIPTION_MAX = 500
_CATEGORY_MAX = 50
# Numeric(20, 3): seventeen digits before the point.
_AMOUNT_MAX_ADJUSTED = 16

# Colours for the categories an import creates, in this order, skipping any the
# user already has. Distinct from each other and from the seeded set.
IMPORT_COLOURS = (
    "#06b6d4",
    "#ec4899",
    "#84cc16",
    "#f97316",
    "#6366f1",
    "#14b8a6",
    "#eab308",
    "#8b5cf6",
)

_UNREADABLE = "We couldn't open that file. Save it as .xlsx, or as CSV UTF-8, and try again."
_TOO_MANY = (
    f"That file has more than {MAX_IMPORT_ROWS:,} rows. "
    "Split it into smaller files and import them one at a time."
)
_HOW_TO_DATE = "Write dates as 2026-10-24, or 24/10/2026 with the day first."


@dataclass
class ImportRow:
    """One expense the file holds, read and checked."""

    row: int
    spent_on: date
    description: str | None
    # As it will be shown: an existing category's own name, or a new one as
    # the file spells it. None is uncategorised.
    category: str | None
    new_category: bool
    amount: Decimal
    currency: str
    duplicate: bool = False
    category_id: uuid.UUID | None = None


@dataclass
class ImportProblem:
    row: int
    column: str
    message: str


@dataclass
class ReadFile:
    rows: list[ImportRow] = field(default_factory=list)
    problems: list[ImportProblem] = field(default_factory=list)


# --- The template ---------------------------------------------------------------


def _validation(sheet_range: str, prompt: str, error: str, **kwargs: Any) -> DataValidation:
    """A column's check, with a hint shown when one of its cells is selected."""
    check = DataValidation(
        allow_blank=True,
        showInputMessage=True,
        prompt=prompt,
        showErrorMessage=True,
        error=error,
        **kwargs,
    )
    check.add(sheet_range)
    return check


def build_template(
    categories: Sequence[Category], currencies: Sequence[str], *, default_currency: str
) -> bytes:
    """The import template as .xlsx bytes, made for one user.

    Empty on purpose: an example row left in by mistake would be imported as an
    expense. Examples live on the instructions sheet instead.
    """
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Expenses"
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = _BOLD
    sheet.freeze_panes = "A2"
    for column, width in enumerate([12, 40, 18, 14, 10], start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    last = _TEMPLATE_ROWS + 1
    for row in range(2, last + 1):
        sheet.cell(row=row, column=1).number_format = _DATE_FORMAT

    names = sorted((category.name for category in categories), key=str.casefold)
    lists = book.create_sheet("Lists")
    lists.sheet_state = "hidden"
    lists.append(["Categories", "Currencies"])
    for row, name in enumerate(names, start=2):
        _text(lists.cell(row=row, column=1), name)
    for row, code in enumerate(currencies, start=2):
        _text(lists.cell(row=row, column=2), code)

    checks = [
        _validation(
            f"A2:A{last}",
            "The day you paid, like 2026-10-24.",
            "Enter a date, like 2026-10-24.",
            type="date",
            operator="greaterThanOrEqual",
            formula1="1",
        ),
        _validation(
            f"D2:D{last}",
            "What you paid, like 12.50. No currency symbol.",
            "Enter an amount more than 0, like 12.50.",
            type="decimal",
            operator="greaterThan",
            formula1="0",
        ),
        _validation(
            f"E2:E{last}",
            f"A code from the list. Leave blank for {default_currency}.",
            f"Pick a currency from the list, or leave it blank for {default_currency}.",
            type="list",
            formula1=f"=Lists!$B$2:$B${len(currencies) + 1}",
        ),
    ]
    if names:
        checks.append(
            _validation(
                f"C2:C{last}",
                "One of your categories. A new name is created when you import.",
                "That isn't one of your categories yet. Keep it to create it when you import.",
                type="list",
                formula1=f"=Lists!$A$2:$A${len(names) + 1}",
                # A warning, not a stop: a new name is allowed, just not by accident.
                errorStyle="warning",
            )
        )
    for check in checks:
        sheet.add_data_validation(check)

    _instructions(book.create_sheet("How to fill this in", 1), default_currency)
    book.active = 0

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _instructions(sheet: Any, default_currency: str) -> None:
    rows: list[list[str]] = [
        ["How to fill this in"],
        [],
        ["Column", "Needed?", "What to put"],
        [_DATE, "Yes", f"The day you paid. {_HOW_TO_DATE}"],
        [_DESCRIPTION, "No", f"What it was, up to {_DESCRIPTION_MAX} characters."],
        [
            _CATEGORY,
            "No",
            "One of your categories, from the list. A name you don't have yet is created "
            "when you import. Leave it blank for Uncategorised.",
        ],
        [_AMOUNT, "Yes", "What you paid, more than 0, like 12.50. Use a dot for decimals."],
        [
            _CURRENCY,
            "No",
            f"A three-letter code from the list, like {default_currency}. "
            f"Leave it blank for {default_currency}, your preferred currency.",
        ],
        [],
        ["Example", "", "2026-10-24 · Lunch with Sam · Food · 18.50 · (blank)"],
        [],
        ["Good to know"],
        ["", "", "Each row on the Expenses sheet is one expense. Empty rows are skipped."],
        ["", "", "Before anything is saved you see what will be imported, and what won't be."],
        [
            "",
            "",
            "A row with the same date, amount, currency and description as an expense you "
            "already have is left out, unless you choose to import it too.",
        ],
        ["", "", "A file exported from the Transactions tab imports as it is."],
        ["", "", "You can also save the Expenses sheet as CSV UTF-8 and import that."],
    ]
    for values in rows:
        sheet.append(values)
    for row in (1, 3, 10, 12):
        for cell in sheet[row]:
            cell.font = _BOLD
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = wrap
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 80


async def import_template(db: AsyncSession, user_id: uuid.UUID, *, default_currency: str) -> bytes:
    categories = await list_categories(db, user_id)
    currencies = sorted(await exchange_rates.supported_currencies(db))
    return build_template(categories, currencies, default_currency=default_currency)


# --- Reading a file ---------------------------------------------------------------


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _filled(rows: Iterable[Sequence[Any]]) -> list[tuple[int, Sequence[Any]]]:
    """The rows with anything in them, numbered as the spreadsheet numbers them.

    Refused as soon as there are too many rather than read to the end and cut
    short: a cut-off import would drop expenses without saying so. Blank rows
    are counted for numbering and then let go, so a sheet that is mostly empty
    costs nothing to hold.
    """
    kept: list[tuple[int, Sequence[Any]]] = []
    for number, values in enumerate(rows, start=1):
        if all(map(_blank, values)):
            continue
        kept.append((number, values))
        # The column names, then one row more than an import may hold.
        if len(kept) > MAX_IMPORT_ROWS + 1:
            raise ImportFileError(_TOO_MANY)
    return kept


def _xlsx_rows(content: bytes) -> list[tuple[int, Sequence[Any]]]:
    """The first worksheet's filled rows, as values (cached results for formulas)."""
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
        if unpacked > _MAX_UNPACKED_BYTES:
            raise ImportFileError(_UNREADABLE)
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except ImportFileError:
        raise
    except Exception as exc:
        # openpyxl fails a broken file in many ways (BadZipFile, KeyError,
        # InvalidFileException, XML errors); to the person they are one problem.
        raise ImportFileError(_UNREADABLE) from exc
    try:
        if not book.worksheets:
            raise ImportFileError(_UNREADABLE)
        return _filled(book.worksheets[0].iter_rows(values_only=True))
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError(_UNREADABLE) from exc
    finally:
        book.close()


def _csv_rows(content: bytes) -> list[tuple[int, Sequence[Any]]]:
    """A CSV file's filled records. UTF-8, with or without the BOM Excel writes."""
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportFileError(_UNREADABLE) from exc
    if "\x00" in text:
        # Not text at all — an image, a PDF.
        raise ImportFileError(_UNREADABLE)
    # The delimiter is whichever the first line uses most: Excel writes ";" in
    # locales where "," is the decimal point. The first line is the column
    # names, which have no commas of their own to throw the count.
    first = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(",;\t", key=first.count)
    try:
        return _filled(csv.reader(StringIO(text, newline=""), delimiter=delimiter))
    except csv.Error as exc:
        raise ImportFileError(_UNREADABLE) from exc


def _sheet_rows(content: bytes) -> list[tuple[int, Sequence[Any]]]:
    if content.startswith(b"PK\x03\x04"):
        return _xlsx_rows(content)
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        raise ImportFileError(
            "That's an older .xls file. Save it as .xlsx or CSV, then import that."
        )
    return _csv_rows(content)


def _shown(value: Any) -> str:
    """A cell as it might look on screen, for quoting in a message."""
    text = str(value).strip()
    return text if len(text) <= 40 else f"{text[:39]}…"


def read_table(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    """The data rows of a file, numbered as the spreadsheet numbers them, each as
    {column: cell} for the columns the importer reads.

    The column names are the first row with anything in it. Other columns are
    ignored — the export's converted-amount column among them. Rows with
    nothing in the columns that matter are skipped.
    """
    rows = _sheet_rows(content)
    if not rows:
        raise ImportFileError("That file is empty. Fill in the template and try again.")

    columns: dict[str, int] = {}
    for index, value in enumerate(rows[0][1]):
        name = None if _blank(value) else _BY_KEY.get(str(value).strip().casefold())
        if name is not None and name not in columns:
            columns[name] = index
    missing = [name for name in _REQUIRED if name not in columns]
    if missing:
        raise ImportFileError(
            f"The first row needs a column named {' and '.join(missing)}. "
            "Download the template to see the layout."
        )

    table: list[tuple[int, dict[str, Any]]] = []
    for number, values in rows[1:]:
        cells = {
            name: values[index] if index < len(values) else None for name, index in columns.items()
        }
        if not all(map(_blank, cells.values())):
            table.append((number, cells))
    return table


# --- Reading a cell ------------------------------------------------------------------

_ISO_DATE = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[T ][\d:.]+)?")
_DAY_FIRST = re.compile(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})")
_PLAIN_NUMBER = re.compile(r"-?\d+(?:\.\d+)?")
_GROUPED_NUMBER = re.compile(r"-?\d{1,3}(?:,\d{3})+(?:\.\d+)?")


def parse_date(value: Any) -> date:
    """A date cell, or text in one of two shapes: ISO, or day first.

    Day first because that is how the people using this app write dates; the
    ISO form is what the template shows. Month-first text is not guessed at:
    05/10 read the wrong way round is a wrong expense nobody notices.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if match := _ISO_DATE.fullmatch(text):
            year, month, day = match.groups()
        elif match := _DAY_FIRST.fullmatch(text):
            day, month, year = match.groups()
        else:
            raise ValueError(text)
        return date(int(year), int(month), int(day))
    raise ValueError(value)


def parse_amount(value: Any) -> Decimal:
    """A number cell, or text like "12.50", "1,250.00" or "250000".

    A float from Excel goes through its shortest repr, so 12.1 arrives as 12.1
    rather than its binary expansion. Commas are only taken as thousands
    separators where they group digits in threes: "12,50" is refused rather
    than read as twelve hundred and fifty.
    """
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError(value)
        return Decimal(repr(value))
    if isinstance(value, Decimal):
        return value
    if isinstance(value, str):
        text = value.strip().replace("\N{NO-BREAK SPACE}", "").replace(" ", "")
        if _PLAIN_NUMBER.fullmatch(text) or _GROUPED_NUMBER.fullmatch(text):
            try:
                return Decimal(text.replace(",", ""))
            except InvalidOperation as exc:
                raise ValueError(text) from exc
    raise ValueError(value)


def _amount_problem(amount: Decimal, currency: str) -> str | None:
    if amount <= 0:
        return "Amount must be more than 0."
    if amount.adjusted() > _AMOUNT_MAX_ADJUSTED:
        return "Amount is too large."
    try:
        check_amount_scale(amount, currency)
    except ValueError:
        places = minor_units(currency)
        return (
            f"Amount must have no more than {places} decimal place{'' if places == 1 else 's'} "
            f"for {currency}."
        )
    return None


class _Categories:
    """Category names in the file, matched to the user's own however they are
    capitalised, and remembered when they are new."""

    def __init__(self, existing: Iterable[Category]) -> None:
        self._exact: dict[str, Category] = {}
        self._folded: dict[str, Category] = {}
        for category in existing:
            self._exact.setdefault(category.name, category)
            self._folded.setdefault(category.name.casefold(), category)
        self._new: dict[str, str] = {}

    def resolve(self, name: str) -> tuple[str, Category | None]:
        """The name to show, and the category if it exists."""
        found = self._exact.get(name) or self._folded.get(name.casefold())
        if found is not None:
            return found.name, found
        # A new name keeps the spelling it first appeared with.
        return self._new.setdefault(name.casefold(), name), None


def _read_row(
    number: int,
    cells: dict[str, Any],
    names: _Categories,
    *,
    supported: frozenset[str],
    default_currency: str,
) -> ImportRow | list[ImportProblem]:
    """One row as an expense, or everything wrong with it — never half of each."""
    problems: list[ImportProblem] = []

    def problem(column: str, message: str) -> None:
        problems.append(ImportProblem(row=number, column=column, message=message))

    spent_on: date | None = None
    raw_date = cells.get(_DATE)
    if _blank(raw_date):
        problem(_DATE, "Enter the date.")
    else:
        try:
            spent_on = parse_date(raw_date)
        except ValueError:
            problem(_DATE, f"“{_shown(raw_date)}” isn’t a date. {_HOW_TO_DATE}")

    raw_currency = cells.get(_CURRENCY)
    currency = default_currency if _blank(raw_currency) else str(raw_currency).strip().upper()
    if currency not in supported:
        problem(
            _CURRENCY,
            f"{currency} can’t be used for new expenses. Pick another currency."
            if _blank(raw_currency)
            else f"“{_shown(raw_currency)}” isn’t a currency you can use. "
            "Use a three-letter code, like SGD.",
        )

    amount: Decimal | None = None
    raw_amount = cells.get(_AMOUNT)
    if _blank(raw_amount):
        problem(_AMOUNT, "Enter the amount.")
    else:
        try:
            amount = parse_amount(raw_amount)
        except ValueError:
            problem(_AMOUNT, f"“{_shown(raw_amount)}” isn’t an amount. Write it like 12.50.")
        else:
            if found := _amount_problem(amount, currency):
                problem(_AMOUNT, found)

    raw_description = cells.get(_DESCRIPTION)
    description = None if _blank(raw_description) else str(raw_description).strip()
    if description is not None and len(description) > _DESCRIPTION_MAX:
        problem(_DESCRIPTION, f"Description must be {_DESCRIPTION_MAX} characters or less.")

    raw_category = cells.get(_CATEGORY)
    category_name: str | None = None
    existing: Category | None = None
    if not _blank(raw_category):
        typed = str(raw_category).strip()
        if len(typed) > _CATEGORY_MAX:
            problem(_CATEGORY, f"Category must be {_CATEGORY_MAX} characters or less.")
        elif typed.casefold() not in _UNCATEGORISED:
            category_name, existing = names.resolve(typed)

    if problems or spent_on is None or amount is None:
        return problems
    return ImportRow(
        row=number,
        spent_on=spent_on,
        description=description,
        category=category_name,
        new_category=category_name is not None and existing is None,
        amount=amount,
        currency=currency,
        category_id=existing.id if existing is not None else None,
    )


def read_rows(
    table: Sequence[tuple[int, dict[str, Any]]],
    categories: Iterable[Category],
    *,
    supported: frozenset[str],
    default_currency: str,
) -> ReadFile:
    """Check every row: each becomes an expense or a list of its problems."""
    names = _Categories(categories)
    result = ReadFile()
    for number, cells in table:
        read = _read_row(
            number, cells, names, supported=supported, default_currency=default_currency
        )
        if isinstance(read, ImportRow):
            result.rows.append(read)
        else:
            result.problems.extend(read)
    return result


def _fingerprint(
    spent_on: date, amount: Decimal, currency: str, description: str | None
) -> tuple[date, Decimal, str, str]:
    """What makes two expenses "the same" for an import: day, money and words.
    Not the category — the one in the file may be newer than the stored one."""
    return spent_on, amount.normalize(), currency, (description or "").strip().casefold()


async def _mark_duplicates(db: AsyncSession, user_id: uuid.UUID, rows: list[ImportRow]) -> None:
    """Flag rows that match an expense already recorded, one for one.

    Counted rather than matched as a set: two identical coffees in the file and
    one already saved leaves one of the file's two as new. That is also what
    makes re-importing a fixed file safe after importing part of it.
    """
    if not rows:
        return
    existing = await db.execute(
        select(Expense.spent_on, Expense.amount, Expense.currency, Expense.description).where(
            Expense.user_id == user_id,
            Expense.spent_on >= min(row.spent_on for row in rows),
            Expense.spent_on <= max(row.spent_on for row in rows),
        )
    )
    remaining = Counter(_fingerprint(*expense) for expense in existing)
    for row in rows:
        key = _fingerprint(row.spent_on, row.amount, row.currency, row.description)
        if remaining[key] > 0:
            row.duplicate = True
            remaining[key] -= 1


async def _read(
    db: AsyncSession, user_id: uuid.UUID, content: bytes, *, default_currency: str
) -> ReadFile:
    table = read_table(content)
    result = read_rows(
        table,
        await list_categories(db, user_id),
        supported=await exchange_rates.supported_currencies(db),
        default_currency=default_currency,
    )
    await _mark_duplicates(db, user_id, result.rows)
    return result


async def preview_import(
    db: AsyncSession, user_id: uuid.UUID, content: bytes, *, default_currency: str
) -> dict[str, Any]:
    """What importing ``content`` would do. Saves nothing."""
    result = await _read(db, user_id, content, default_currency=default_currency)
    return {
        "rows": [asdict(row) for row in result.rows],
        "problems": [asdict(problem) for problem in result.problems],
    }


# --- Saving ---------------------------------------------------------------------


def _colours(taken: Iterable[str]) -> Iterable[str]:
    """The import palette, unused colours first, then round again."""
    used = {colour.lower() for colour in taken}
    fresh = [colour for colour in IMPORT_COLOURS if colour not in used]
    while True:
        yield from fresh or IMPORT_COLOURS
        fresh = list(IMPORT_COLOURS)


async def import_expenses(
    db: AsyncSession,
    user_id: uuid.UUID,
    content: bytes,
    *,
    default_currency: str,
    include_duplicates: bool,
) -> dict[str, Any]:
    """Save the ready rows of ``content``, and the categories they name that
    don't exist yet, in one transaction. Rows with problems are left out —
    the preview listed them — and so are likely duplicates unless
    ``include_duplicates``."""
    result = await _read(db, user_id, content, default_currency=default_currency)
    chosen = [row for row in result.rows if include_duplicates or not row.duplicate]

    existing = await list_categories(db, user_id)
    palette = iter(_colours(category.colour for category in existing))
    created: dict[str, Category] = {}
    for row in chosen:
        if row.new_category and row.category is not None:
            key = row.category.casefold()
            if key not in created:
                # The id is set here so the expenses below can point at it.
                created[key] = Category(
                    id=uuid.uuid4(), user_id=user_id, name=row.category, colour=next(palette)
                )

    # Flushed first: the expenses reference these rows, and with no
    # relationship between the two models nothing else orders the inserts.
    # Still one transaction — the commit below is the only one.
    db.add_all(created.values())
    try:
        await db.flush()
    except IntegrityError as exc:
        # A category of the same name was created meanwhile, in another tab.
        await db.rollback()
        raise CategoryNameTakenError(next(iter(created.values())).name) from exc
    db.add_all(
        Expense(
            user_id=user_id,
            category_id=(
                created[row.category.casefold()].id
                if row.new_category and row.category is not None
                else row.category_id
            ),
            amount=row.amount,
            currency=row.currency,
            description=row.description,
            spent_on=row.spent_on,
        )
        for row in chosen
    )
    await db.commit()

    return {
        "imported": len(chosen),
        "duplicates_skipped": len(result.rows) - len(chosen),
        "problem_rows_skipped": len({problem.row for problem in result.problems}),
        "categories_created": [category.name for category in created.values()],
    }
