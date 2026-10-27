"""Reading an import file, cell by cell and row by row, without a database."""

import codecs
import uuid
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.models.finance import Category
from app.services.finance.errors import ImportFileError
from app.services.finance.imports import (
    MAX_IMPORT_ROWS,
    ImportProblem,
    ImportRow,
    build_template,
    parse_amount,
    parse_date,
    read_rows,
    read_table,
)

SUPPORTED = frozenset({"USD", "SGD", "VND"})
FOOD = Category(id=uuid.uuid4(), name="Food", colour="#22c55e")
HEADER = "Date,Description,Category,Amount,Currency\n"


def csv_bytes(*lines: str, header: str = HEADER, bom: bool = False) -> bytes:
    text = header + "\n".join(lines) + "\n"
    return (codecs.BOM_UTF8 if bom else b"") + text.encode()


def xlsx_bytes(rows: list[list[Any]]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for number, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            if value is not None:
                sheet.cell(row=number, column=column, value=value)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def read(
    content: bytes, default_currency: str = "SGD"
) -> tuple[list[ImportRow], list[ImportProblem]]:
    result = read_rows(
        read_table(content), [FOOD], supported=SUPPORTED, default_currency=default_currency
    )
    return result.rows, result.problems


# --- Cells ------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (datetime(2026, 10, 24, 0, 0), date(2026, 10, 24)),
        (date(2026, 10, 24), date(2026, 10, 24)),
        ("2026-10-24", date(2026, 10, 24)),
        (" 2026-10-24 ", date(2026, 10, 24)),
        ("2026-10-24 00:00:00", date(2026, 10, 24)),
        ("2026-1-5", date(2026, 1, 5)),
        ("24/10/2026", date(2026, 10, 24)),
        ("5/1/2026", date(2026, 1, 5)),
        ("24.10.2026", date(2026, 10, 24)),
    ],
)
def test_dates_are_read_as_iso_or_day_first(value: Any, expected: date) -> None:
    assert parse_date(value) == expected


@pytest.mark.parametrize(
    "value", ["31/02/2026", "10/24/2026", "Oct 24 2026", "24/10/26", "", 45000, 1.5, None]
)
def test_anything_else_is_not_a_date(value: Any) -> None:
    with pytest.raises(ValueError):
        parse_date(value)


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (12, Decimal("12")),
        (12.1, Decimal("12.1")),
        (Decimal("12.50"), Decimal("12.50")),
        ("12.50", Decimal("12.50")),
        ("1,250.50", Decimal("1250.50")),
        ("250,000", Decimal("250000")),
        ("250 000", Decimal("250000")),
        # How Excel and some banks group thousands in European locales.
        ("250\N{NO-BREAK SPACE}000", Decimal("250000")),
        ("-5", Decimal("-5")),
    ],
)
def test_amounts_keep_the_decimals_written(value: Any, expected: Decimal) -> None:
    assert parse_amount(value) == expected
    assert str(parse_amount(value)) == str(expected)


@pytest.mark.parametrize(
    "value", ["12,50", "1,25,000", "$12.50", "12.5.0", "abc", True, float("nan")]
)
def test_anything_else_is_not_an_amount(value: Any) -> None:
    with pytest.raises(ValueError):
        parse_amount(value)


# --- Tables ------------------------------------------------------------------------


def test_rows_are_numbered_as_the_spreadsheet_numbers_them() -> None:
    content = xlsx_bytes(
        [
            [],
            ["Date", "Amount"],
            ["2026-10-01", 5],
            [],
            [None, None],
            ["2026-10-02", 6],
        ]
    )

    assert [number for number, _ in read_table(content)] == [3, 6]


def test_column_names_are_matched_loosely_and_extra_columns_ignored() -> None:
    # The export's converted column ("Amount (SGD)") must not be taken for Amount.
    content = csv_bytes("9.5,2026-10-01,ignored", header=" AMOUNT ,date,Amount (SGD)\n")

    assert read_table(content) == [(2, {"Amount": "9.5", "Date": "2026-10-01"})]


@pytest.mark.parametrize("delimiter", [";", "\t"])
def test_a_csv_in_another_delimiter_is_read(delimiter: str) -> None:
    content = f"Date{delimiter}Amount\n2026-10-01{delimiter}12,50\n".encode()

    [(_, cells)] = read_table(content)

    assert cells == {"Date": "2026-10-01", "Amount": "12,50"}


def test_excels_utf8_bom_is_not_part_of_the_first_column_name() -> None:
    [(_, cells)] = read_table(csv_bytes("2026-10-01,Phở,,85000,VND", bom=True))

    assert cells["Date"] == "2026-10-01"
    assert cells["Description"] == "Phở"


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"", "That file is empty"),
        (b"\n\n", "That file is empty"),
        (b"Date,Price\n2026-10-01,5\n", "needs a column named Amount"),
        (b"Name\nx\n", "needs a column named Date and Amount"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "older .xls file"),
        (b"PK\x03\x04not really a zip", "couldn't open that file"),
        (b"\x89PNG\r\n\x1a\n\x00\x00", "couldn't open that file"),
        ("Date,Amount\n".encode("utf-16"), "couldn't open that file"),
    ],
)
def test_a_file_that_cant_be_read_says_why(content: bytes, message: str) -> None:
    with pytest.raises(ImportFileError, match=message):
        read_table(content)


def test_too_many_rows_are_refused_rather_than_cut_short() -> None:
    lines = [f"2026-10-01,{n}" for n in range(1, MAX_IMPORT_ROWS + 2)]

    with pytest.raises(ImportFileError, match="more than 5,000 rows"):
        read_table(csv_bytes(*lines, header="Date,Amount\n"))
    assert len(read_table(csv_bytes(*lines[:-1], header="Date,Amount\n"))) == MAX_IMPORT_ROWS


def test_the_template_reads_back_as_an_empty_import() -> None:
    content = build_template([FOOD], ["SGD", "USD"], default_currency="SGD")

    assert read_table(content) == []
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ["Expenses", "How to fill this in", "Lists"]
    assert book["Lists"].sheet_state == "hidden"
    assert [c.value for c in book["Expenses"][1]] == [
        "Date",
        "Description",
        "Category",
        "Amount",
        "Currency",
    ]
    lists = list(book["Lists"].iter_rows(values_only=True))
    assert lists == [("Categories", "Currencies"), ("Food", "SGD"), (None, "USD")]
    ranges = {str(v.sqref): v.type for v in book["Expenses"].data_validations.dataValidation}
    assert ranges == {
        "A2:A1001": "date",
        "D2:D1001": "decimal",
        "E2:E1001": "list",
        "C2:C1001": "list",
    }


# --- Rows --------------------------------------------------------------------------


def test_a_good_row_becomes_an_expense() -> None:
    [row], problems = read(csv_bytes("24/10/2026, Lunch ,food,18.50,usd"))

    assert problems == []
    assert row == ImportRow(
        row=2,
        spent_on=date(2026, 10, 24),
        description="Lunch",
        # Matched however it was typed, and shown as the category is named.
        category="Food",
        new_category=False,
        amount=Decimal("18.50"),
        currency="USD",
        category_id=FOOD.id,
    )


def test_blank_currency_is_the_preferred_one_and_blank_category_uncategorised() -> None:
    [row], _ = read(csv_bytes("2026-10-24,,,85000,"), default_currency="VND")
    [named], _ = read(csv_bytes("2026-10-24,,Uncategorized,1,"))

    assert (row.currency, row.category, row.description) == ("VND", None, None)
    assert named.category is None


def test_a_new_category_keeps_the_spelling_it_first_appears_with() -> None:
    rows, _ = read(csv_bytes("2026-10-24,,Pets,5,", "2026-10-25,,PETS,6,"))

    assert [(r.category, r.new_category, r.category_id) for r in rows] == [
        ("Pets", True, None),
        ("Pets", True, None),
    ]


def test_every_problem_in_a_row_is_reported_and_the_row_kept_out() -> None:
    rows, problems = read(
        csv_bytes(
            "31/02/2026,,,12.345,USD",
            ",x,,,",
            "2026-10-24,,,-3,",
            "2026-10-24,,,1.5,VND",
            "2026-10-24,,,5,XYZ",
            "2026-10-24," + "d" * 501 + ",,5,",
            "2026-10-24,,," + "9" * 18 + ",",
        )
    )

    assert rows == []
    how_to = "Write dates as 2026-10-24, or 24/10/2026 with the day first."
    assert [(p.row, p.column, p.message) for p in problems] == [
        (2, "Date", f"“31/02/2026” isn’t a date. {how_to}"),
        (2, "Amount", "Amount must have no more than 2 decimal places for USD."),
        (3, "Date", "Enter the date."),
        (3, "Amount", "Enter the amount."),
        (4, "Amount", "Amount must be more than 0."),
        (5, "Amount", "Amount must have no more than 0 decimal places for VND."),
        (6, "Currency", "“XYZ” isn’t a currency you can use. Use a three-letter code, like SGD."),
        (7, "Description", "Description must be 500 characters or less."),
        (8, "Amount", "Amount is too large."),
    ]


def test_a_preferred_currency_that_cant_be_used_is_reported_on_blank_rows() -> None:
    _, problems = read(csv_bytes("2026-10-24,,,5,"), default_currency="AMD")

    assert [p.message for p in problems] == [
        "AMD can’t be used for new expenses. Pick another currency."
    ]


def test_excel_numbers_and_dates_arrive_as_typed_cells() -> None:
    content = xlsx_bytes(
        [
            ["Date", "Description", "Category", "Amount", "Currency"],
            [datetime(2026, 10, 24), 7, "Food", 12.1, "SGD"],
        ]
    )

    [row], problems = read(content)

    assert problems == []
    assert (row.spent_on, row.description, row.amount) == (date(2026, 10, 24), "7", Decimal("12.1"))
