"""The Transactions export: POST /finance/expenses/export, answered with an .xlsx file."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker
from tests.support.finance_api import make_category, make_expense
from tests.support.ids import OTHER_USER_ID

from app.models.finance import Expense
from app.services import exchange_rates

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
URL = "/api/v1/finance/expenses/export"


async def _book(client: AsyncClient, ids: list[str], filters: list[str] | None = None) -> Workbook:
    resp = await client.post(URL, json={"expense_ids": ids, "filters": filters or []})
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == XLSX
    assert resp.headers["content-disposition"].startswith('attachment; filename="transactions-')
    book = load_workbook(BytesIO(resp.content))
    assert book.sheetnames == ["Transactions", "About this export"]
    return book


async def _rows(client: AsyncClient, ids: list[str]) -> list[tuple[Any, ...]]:
    return list((await _book(client, ids))["Transactions"].iter_rows(values_only=True))


def _dec(value: Any) -> Decimal:
    return Decimal(str(value))


async def test_one_row_per_expense_newest_first_with_real_dates_and_numbers(
    finance_client: AsyncClient, stub_rates: None
) -> None:
    food = await make_category(finance_client, name="Food")
    lunch = await make_expense(
        finance_client,
        amount="12.80",
        currency="SGD",
        description="Lunch",
        category_id=food["id"],
        spent_on="2026-08-01",
    )
    later = await make_expense(
        finance_client, amount="5.00", description=None, spent_on="2026-08-20"
    )

    # Sent oldest first: the file still reads newest first, like the screen.
    rows = await _rows(finance_client, [lunch["id"], later["id"]])

    assert rows[0] == ("Date", "Description", "Category", "Amount", "Currency", "Amount (USD)")
    assert rows[1][0].date() == date(2026, 8, 20)
    assert rows[1][1:3] == (None, "Uncategorised")
    assert _dec(rows[1][3]) == Decimal("5")
    assert rows[2][1:3] == ("Lunch", "Food")
    # 12.80 SGD at 1.28 per USD.
    assert (_dec(rows[2][3]), rows[2][4], _dec(rows[2][5])) == (
        Decimal("12.8"),
        "SGD",
        Decimal("10"),
    )
    assert len(rows) == 3


async def test_only_the_listed_expenses_are_exported(
    finance_client: AsyncClient, stub_rates: None
) -> None:
    taxi = await make_expense(finance_client, description="Taxi")
    await make_expense(finance_client, description="Rent")

    rows = await _rows(finance_client, [taxi["id"]])
    assert [r[1] for r in rows[1:]] == ["Taxi"]


async def test_someone_elses_ids_and_unknown_ids_are_left_out(
    finance_client: AsyncClient,
    stub_rates: None,
    db_factory: async_sessionmaker[AsyncSession],
) -> None:
    mine = await make_expense(finance_client, description="Mine")
    async with db_factory() as session:
        theirs = Expense(
            user_id=OTHER_USER_ID,
            amount=Decimal("3.00"),
            currency="USD",
            description="Theirs",
            spent_on=date(2026, 8, 19),
        )
        session.add(theirs)
        await session.commit()

    unknown = "00000000-0000-0000-0000-000000000000"
    rows = await _rows(finance_client, [mine["id"], str(theirs.id), unknown])
    assert [r[1] for r in rows[1:]] == ["Mine"]


async def test_a_description_that_looks_like_a_formula_stays_text(
    finance_client: AsyncClient, stub_rates: None
) -> None:
    sneaky = await make_expense(finance_client, description='=HYPERLINK("http://evil.test","x")')

    sheet = (await _book(finance_client, [sneaky["id"]]))["Transactions"]

    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == '=HYPERLINK("http://evil.test","x")'


async def test_a_missing_rate_blanks_the_converted_cell_not_the_export(
    finance_client: AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def _usd_only(_db: object) -> dict[str, Decimal]:
        return {"USD": Decimal("1")}

    sgd = await make_expense(finance_client, amount="12.80", currency="SGD")
    monkeypatch.setattr(exchange_rates, "get_rates", _usd_only)

    rows = await _rows(finance_client, [sgd["id"]])
    assert rows[1][4] == "SGD"
    assert rows[1][5] is None


async def test_the_about_sheet_lists_the_filters_as_plain_text(
    finance_client: AsyncClient, stub_rates: None
) -> None:
    expense = await make_expense(finance_client)

    book = await _book(finance_client, [expense["id"]], ["Search: taxi", "=1+1"])
    about = book["About this export"]
    lines = [row[0] for row in about.iter_rows(values_only=True) if row]

    assert "Search: taxi" in lines
    assert "=1+1" in lines
    assert all(cell.data_type == "s" for cell in about["A"] if cell.value == "=1+1")


async def test_no_filters_says_so(finance_client: AsyncClient, stub_rates: None) -> None:
    book = await _book(finance_client, [])
    lines = [row[0] for row in book["About this export"].iter_rows(values_only=True) if row]
    assert "None — every transaction" in lines


async def test_an_empty_export_is_still_a_valid_file(
    finance_client: AsyncClient, stub_rates: None
) -> None:
    assert len(await _rows(finance_client, [])) == 1


async def test_ids_must_be_ids(finance_client: AsyncClient) -> None:
    resp = await finance_client.post(URL, json={"expense_ids": ["not-an-id"]})
    assert resp.status_code == 422
