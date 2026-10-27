"""Importing expenses: the template, the preview, and the import itself."""

from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from httpx import AsyncClient, Response
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker
from tests.support.finance_api import make_category, make_expense
from tests.support.ids import OTHER_USER_ID

from app.models.finance import Category
from app.services.finance.imports import MAX_IMPORT_BYTES

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BASE = "/api/v1/finance/expenses"
HEADER = "Date,Description,Category,Amount,Currency\n"


def csv_file(*lines: str) -> bytes:
    return (HEADER + "\n".join(lines) + "\n").encode()


async def upload(client: AsyncClient, path: str, content: bytes, **params: Any) -> Response:
    return await client.post(
        f"{BASE}/{path}",
        content=content,
        params=params,
        headers={"Content-Type": "application/octet-stream"},
    )


async def preview(client: AsyncClient, content: bytes) -> dict[str, Any]:
    resp = await upload(client, "import/preview", content)
    assert resp.status_code == 200, resp.text
    return dict(resp.json())


async def expenses(client: AsyncClient) -> list[dict[str, Any]]:
    return list((await client.get(BASE)).json())


async def test_the_template_carries_the_users_categories(
    finance_client: AsyncClient, db_factory: async_sessionmaker[AsyncSession]
) -> None:
    await make_category(finance_client, name="Food")
    async with db_factory() as session:
        session.add(Category(user_id=OTHER_USER_ID, name="Not mine"))
        await session.commit()

    resp = await finance_client.get(f"{BASE}/import/template")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == XLSX
    assert resp.headers["content-disposition"] == 'attachment; filename="expenses-template.xlsx"'
    book = load_workbook(BytesIO(resp.content))
    categories = [row[0] for row in book["Lists"].iter_rows(min_row=2, values_only=True) if row[0]]
    assert categories == ["Food"]
    currencies = [row[1] for row in book["Lists"].iter_rows(min_row=2, values_only=True) if row[1]]
    assert {"USD", "SGD", "VND"} <= set(currencies)


async def test_a_preview_says_what_would_happen_and_saves_nothing(
    finance_client: AsyncClient,
) -> None:
    await make_category(finance_client, name="Food")

    result = await preview(
        finance_client,
        csv_file(
            "2026-10-24,Lunch,Food,18.50,SGD",
            "24/10/2026,Vet,Pets,80,",
            "2026-10-25,Broken,,abc,",
        ),
    )

    assert result["rows"] == [
        {
            "row": 2,
            "spent_on": "2026-10-24",
            "description": "Lunch",
            "category": "Food",
            "new_category": False,
            "amount": "18.50",
            "currency": "SGD",
            "duplicate": False,
        },
        {
            "row": 3,
            "spent_on": "2026-10-24",
            "description": "Vet",
            "category": "Pets",
            "new_category": True,
            "amount": "80",
            # Blank: the profile's preferred currency.
            "currency": "USD",
            "duplicate": False,
        },
    ]
    assert result["problems"] == [
        {"row": 4, "column": "Amount", "message": "“abc” isn’t an amount. Write it like 12.50."}
    ]
    assert await expenses(finance_client) == []


async def test_an_import_saves_the_ready_rows_and_creates_their_categories(
    finance_client: AsyncClient,
) -> None:
    await make_category(finance_client, name="Food")
    content = csv_file(
        "2026-10-24,Lunch,Food,18.50,SGD",
        "2026-10-24,Vet,Pets,80,",
        "2026-10-25,Treats,pets,4.20,",
        "2026-10-25,Broken,,abc,",
    )

    resp = await upload(finance_client, "import", content)

    assert resp.status_code == 201, resp.text
    assert resp.json() == {
        "imported": 3,
        "duplicates_skipped": 0,
        "problem_rows_skipped": 1,
        "categories_created": ["Pets"],
    }
    categories = {
        c["name"]: c for c in (await finance_client.get("/api/v1/finance/categories")).json()
    }
    pets = categories["Pets"]
    assert pets["colour"].startswith("#")
    saved = sorted(
        (e["description"], e["category_id"], Decimal(e["amount"]), e["currency"])
        for e in await expenses(finance_client)
    )
    assert saved == [
        ("Lunch", categories["Food"]["id"], Decimal("18.50"), "SGD"),
        ("Treats", pets["id"], Decimal("4.20"), "USD"),
        ("Vet", pets["id"], Decimal("80"), "USD"),
    ]


async def test_rows_already_recorded_are_left_out_unless_asked(
    finance_client: AsyncClient,
) -> None:
    await make_expense(
        finance_client, amount="4.50", currency="SGD", description="Coffee", spent_on="2026-10-24"
    )
    # Two coffees in the file, one already saved: one of the two is new.
    content = csv_file("2026-10-24,coffee,,4.5,SGD", "2026-10-24,Coffee,,4.50,SGD")

    rows = (await preview(finance_client, content))["rows"]
    skipped = await upload(finance_client, "import", content)
    kept = await upload(finance_client, "import", content, include_duplicates="true")

    assert [r["duplicate"] for r in rows] == [True, False]
    assert skipped.json()["imported"] == 1 and skipped.json()["duplicates_skipped"] == 1
    # By now two are saved, so both rows of the file match — and are kept anyway.
    assert kept.json()["imported"] == 2 and kept.json()["duplicates_skipped"] == 0
    assert len(await expenses(finance_client)) == 4


async def test_an_exported_file_imports_as_it_is(
    finance_client: AsyncClient, stub_rates: None
) -> None:
    food = await make_category(finance_client, name="Groceries")
    first = await make_expense(
        finance_client,
        amount="250000",
        currency="VND",
        description="Market",
        category_id=food["id"],
        spent_on="2026-10-02",
    )
    second = await make_expense(finance_client, amount="9.99", description=None)
    exported = await finance_client.post(
        f"{BASE}/export", json={"expense_ids": [first["id"], second["id"]]}
    )

    result = await preview(finance_client, exported.content)

    assert result["problems"] == []
    assert [
        (r["spent_on"], r["category"], r["amount"], r["duplicate"]) for r in result["rows"]
    ] == [
        ("2026-10-02", "Groceries", "250000", True),
        ("2026-08-19", None, "9.99", True),
    ]


async def test_a_file_that_cant_be_read_is_a_422(finance_client: AsyncClient) -> None:
    resp = await upload(finance_client, "import/preview", b"\xd0\xcf\x11\xe0 an old xls")

    assert resp.status_code == 422
    assert resp.json()["error"] == {
        "code": "IMPORT_FILE_REJECTED",
        "message": "That's an older .xls file. Save it as .xlsx or CSV, then import that.",
    }


@pytest.mark.parametrize("path", ["import/preview", "import"])
async def test_an_oversized_file_is_refused(finance_client: AsyncClient, path: str) -> None:
    resp = await upload(finance_client, path, b"x" * (MAX_IMPORT_BYTES + 1))

    assert resp.status_code == 413
    assert resp.json()["error"]["code"] == "IMPORT_FILE_TOO_LARGE"
    assert await expenses(finance_client) == []
